"""Test documents, built at test time rather than committed as binaries.

The text-layer-free PDF matters most: FR-2's rejection is the behaviour most likely to
regress silently, and a scanned PDF in the repository would be a binary nobody can
review. Building one from primitives keeps it inspectable — it is a page with drawing
operators and no text operators, which is exactly what a scan is.
"""

from __future__ import annotations

import io
import zlib


def _pdf(objects: list[bytes]) -> bytes:
    """Assemble a minimal but valid PDF with a correct xref table."""
    out = io.BytesIO()
    out.write(b"%PDF-1.7\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")

    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    out.write(trailer.encode())
    return out.getvalue()


def _content_stream(data: bytes) -> bytes:
    compressed = zlib.compress(data)
    return (
        b"<< /Length "
        + str(len(compressed)).encode()
        + b" /Filter /FlateDecode >>\nstream\n"
        + compressed
        + b"\nendstream"
    )


def text_pdf(pages: list[str]) -> bytes:
    """A PDF whose pages carry a real text layer, one `Tj` per line."""
    page_ids = [4 + i * 2 for i in range(len(pages))]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{pid} 0 R".encode() for pid in page_ids)
        + f"] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, body in enumerate(pages):
        content = [b"BT\n/F1 12 Tf\n72 720 Td\n14 TL"]
        for line in body.split("\n"):
            escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            content.append(b"(" + escaped.encode("latin-1", "replace") + b") Tj T*")
        content.append(b"ET")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {page_ids[index] + 1} 0 R >>".encode()
        )
        objects.append(_content_stream(b"\n".join(content)))
    return _pdf(objects)


def scanned_pdf(page_count: int = 2) -> bytes:
    """A PDF with no text operators at all — what a scan looks like to a parser.

    Each page draws a filled rectangle, standing in for the scanned image. There is
    nothing for `extract_text` to return, which is the condition FR-2 rejects on.
    """
    page_ids = [3 + i * 2 for i in range(page_count)]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{pid} 0 R".encode() for pid in page_ids)
        + f"] /Count {page_count} >>".encode(),
    ]
    for index in range(page_count):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {page_ids[index] + 1} 0 R >>".encode()
        )
        objects.append(_content_stream(b"0.5 g\n72 72 468 648 re\nf"))
    return _pdf(objects)


def docx_file(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        created = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                created.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def txt_file(text: str) -> bytes:
    return text.encode("utf-8")
