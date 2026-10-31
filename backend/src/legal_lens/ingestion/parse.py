"""File bytes → pages of text.

A PDF has real pages and they are kept as-is. DOCX and TXT have none, so pages are
synthesised at paragraph boundaries against a character budget: a citation has to name
*somewhere*, and "page 3 of the memo" is a location a user can check, whereas a single
10,000-line page is not.

Nothing here rewrites text. See the package docstring for why.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class ParsedPage:
    page: int
    text: str

    @property
    def char_count(self) -> int:
        return len(self.text)


@dataclass(frozen=True)
class ParseResult:
    pages: list[ParsedPage]
    #: Pages the parser reached but which carried no usable text layer. A PDF where
    #: this equals the page count is the FR-2 scan case.
    empty_pages: int

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def total_chars(self) -> int:
        return sum(p.char_count for p in self.pages)

    @property
    def is_scan(self) -> bool:
        """Every page was reached and none carried a usable text layer.

        That pair is the FR-2 condition: not "the file failed to parse", but "it
        parsed and there was nothing in it" — which is what a scan looks like.
        """
        return self.page_count > 0 and self.empty_pages == self.page_count


class UnreadableFile(Exception):
    """The bytes could not be parsed at all — corrupt, encrypted, or mislabelled."""


def parse(
    data: bytes, suffix: str, *, min_chars_per_text_page: int, page_budget_chars: int
) -> ParseResult:
    if suffix == ".pdf":
        return _parse_pdf(data, min_chars_per_text_page=min_chars_per_text_page)
    if suffix == ".docx":
        return _paginate(_docx_paragraphs(data), page_budget_chars)
    if suffix == ".txt":
        return _paginate(_text_paragraphs(data), page_budget_chars)
    raise UnreadableFile(f"unsupported suffix {suffix}")


def _parse_pdf(data: bytes, *, min_chars_per_text_page: int) -> ParseResult:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        # An encrypted PDF may still open; the empty-password case is common enough in
        # legal exports to be worth trying before giving up.
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise UnreadableFile("encrypted pdf") from exc
        raw_pages = list(reader.pages)
    except UnreadableFile:
        raise
    except (PdfReadError, OSError, ValueError) as exc:
        raise UnreadableFile(str(exc)) from exc

    pages: list[ParsedPage] = []
    empty = 0
    for index, source in enumerate(raw_pages, start=1):
        try:
            text = source.extract_text() or ""
        except Exception:
            # One unreadable page should not lose the other four hundred. It counts as
            # empty, which is what decides the scan verdict below.
            log.warning("page extraction failed", extra={"page": index})
            text = ""
        # Whitespace-only text is no text layer: a scanned page commonly yields a
        # handful of stray form-feed characters rather than nothing at all.
        if len(text.strip()) < min_chars_per_text_page:
            empty += 1
        pages.append(ParsedPage(page=index, text=text))
    return ParseResult(pages=pages, empty_pages=empty)


def _docx_paragraphs(data: bytes) -> list[str]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, OSError) as exc:
        raise UnreadableFile(str(exc)) from exc

    blocks = [p.text for p in document.paragraphs]
    # Table cells carry real contract terms — payment schedules and fee tables live in
    # them — so a parser that walked only paragraphs would silently drop the numbers
    # most likely to be asked about.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return blocks


def _text_paragraphs(data: bytes) -> list[str]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding).split("\n")
        except UnicodeDecodeError:
            continue
    raise UnreadableFile("could not decode as text")


def _paginate(blocks: list[str], budget: int) -> ParseResult:
    """Group blocks into pages, breaking only between them.

    Breaking mid-paragraph would put half a sentence on each of two pages, and a
    citation to either would quote text the user cannot find where it says to look.
    """
    pages: list[ParsedPage] = []
    current: list[str] = []
    size = 0

    for block in blocks:
        # `+ 1` for the newline that will rejoin them.
        if current and size + len(block) + 1 > budget:
            pages.append(ParsedPage(page=len(pages) + 1, text="\n".join(current)))
            current, size = [], 0
        current.append(block)
        size += len(block) + 1

    if current:
        pages.append(ParsedPage(page=len(pages) + 1, text="\n".join(current)))

    empty = sum(1 for p in pages if not p.text.strip())
    return ParseResult(pages=pages, empty_pages=empty)
